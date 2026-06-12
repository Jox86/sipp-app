from rest_framework import viewsets, filters
from rest_framework.permissions import AllowAny
from django_filters.rest_framework import DjangoFilterBackend
from apps.catalog.models import Empresa, Producto, Servicio, PedidoExtra
from api.v1.serializers.catalog_serializers import (
    EmpresaSerializer, ProductoSerializer, ServicioSerializer, PedidoExtraSerializer
)
from rest_framework.decorators import action
from rest_framework.response import Response
import openpyxl
from django.core.files.uploadedfile import UploadedFile

class EmpresaViewSet(viewsets.ModelViewSet):
    serializer_class = EmpresaSerializer
    permission_classes = [AllowAny]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    filterset_fields = ['activo']
    search_fields = ['nombre', 'encargado']
    queryset = Empresa.objects.prefetch_related('productos', 'servicios').all()

    @action(detail=False, methods=['post'])
    def upload_products(self, request):
        """Subir productos desde archivo Excel"""
        file = request.FILES.get('file')
        empresa_id = request.data.get('empresa_id')
        
        if not file or not empresa_id:
            return Response({'error': 'Archivo y empresa requeridos'}, status=400)
        
        try:
            empresa = Empresa.objects.get(id=empresa_id)
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            
            productos_creados = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:
                    Producto.objects.create(
                        empresa=empresa,
                        nombre=str(row[0]) if row[0] else '',
                        tipo=str(row[1]) if len(row) > 1 and row[1] else '',
                        precio=float(row[2]) if len(row) > 2 and row[2] else 0,
                        categoria=str(row[3]) if len(row) > 3 and row[3] else '',
                    )
                    productos_creados += 1
            
            return Response({
                'message': f'{productos_creados} productos cargados',
                'count': productos_creados
            })
        except Exception as e:
            return Response({'error': str(e)}, status=500)

    @action(detail=False, methods=['post'])
    def upload_services(self, request):  # <-- MISMA indentación que upload_products
        file = request.FILES.get('file')
        empresa_id = request.data.get('empresa_id')
        
        if not file or not empresa_id:
            return Response({'error': 'Archivo y empresa requeridos'}, status=400)
        
        try:
            empresa = Empresa.objects.get(id=empresa_id)
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            
            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:
                    Servicio.objects.create(
                        empresa=empresa,
                        nombre=str(row[0]),
                        descripcion=str(row[1]) if len(row) > 1 and row[1] else '',
                        precio=float(row[2]) if len(row) > 2 and row[2] else 0,
                    )
                    count += 1
            
            return Response({'message': f'{count} servicios cargados', 'count': count})
        except Exception as e:
            return Response({'error': str(e)}, status=500)
            

class ProductoViewSet(viewsets.ModelViewSet):
    serializer_class = ProductoSerializer
    permission_classes = [AllowAny]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    filterset_fields = ['empresa', 'categoria']
    search_fields = ['nombre']
    queryset = Producto.objects.select_related('empresa').all()


class ServicioViewSet(viewsets.ModelViewSet):
    serializer_class = ServicioSerializer
    permission_classes = [AllowAny]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    filterset_fields = ['empresa']
    search_fields = ['nombre']
    queryset = Servicio.objects.select_related('empresa').all()


class PedidoExtraViewSet(viewsets.ModelViewSet):
    serializer_class = PedidoExtraSerializer
    permission_classes = [AllowAny]
    queryset = PedidoExtra.objects.select_related('usuario', 'proyecto').all()
